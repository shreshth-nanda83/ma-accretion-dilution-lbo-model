"""
Comprehensive Verification Test Suite for M&A Accretion/Dilution & LBO Engine
=============================================================================
Validates all institutional CV metrics:
  - Dataset Ingestion & SQLite database state
  - Flagship Deal EV: ₹4,200 Cr (12.0x on ₹350 Cr Target EBITDA)
  - 5-Year Cumulative Synergies: ₹380 Cr (40, 75, 85, 90, 90)
  - Acquirer Standalone EPS: ₹45.00 -> Year 2 Pro-Forma EPS: ₹47.88 (+6.4% Accretion)
  - Structure #4 Net Debt / EBITDA: 3.08x (strictly below 3.2x covenant)
  - 5-Year Project IRR: 14.2% Base -> 18.7% Optimal Hybrid
  - Excel generation: 6 tabs verified
"""

import os
import sys

WORKSPACE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if WORKSPACE_DIR not in sys.path:
    sys.path.insert(0, WORKSPACE_DIR)

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import sqlite3
import openpyxl
from src.accretion_dilution_engine import AccretionDilutionEngine, FlagshipDealInputs
from src.debt_structures_irr import DebtStructuresOptimizer
from src.kaggle_data_loader import execute_screener_sql, DB_PATH
from src.generate_excel_model import EXCEL_OUTPUT_PATH, generate_wall_street_excel_model


def test_sqlite_database():
    assert os.path.exists(DB_PATH), f"Database not found at {DB_PATH}"
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute("SELECT count(*) FROM fundamentals")
    fund_count = cur.fetchone()[0]
    assert fund_count > 1000, f"Expected >1000 rows in fundamentals, got {fund_count}"

    cur.execute("SELECT count(*) FROM v_latest_financials")
    view_count = cur.fetchone()[0]
    assert view_count > 400, f"Expected >400 tickers in view, got {view_count}"
    conn.close()
    print(f"[TEST PASS] SQLite Database: {fund_count} financial statement rows across {view_count} tickers.")


def test_sql_screener():
    df_screener = execute_screener_sql()
    assert len(df_screener) == 25, f"Expected 25 top screened targets, got {len(df_screener)}"
    assert "ticker" in df_screener.columns
    assert "synergy_opportunity_score" in df_screener.columns
    assert (df_screener["ebitda_margin_pct"] >= 15.0).all(), "All screened targets must have EBITDA margin >= 15%"
    assert (df_screener["net_debt_to_ebitda"] <= 2.50).all(), "All screened targets must have Net Debt/EBITDA <= 2.5x"
    print(f"[TEST PASS] SQL Target Screener: 25 Institutional Targets successfully screened.")


def test_deal_calibration_and_accretion():
    engine = AccretionDilutionEngine()
    su = engine.calculate_sources_and_uses()
    assert su["sources_uses_balanced"], "Sources and Uses must balance exactly"
    assert su["enterprise_value"] == 4200.0, f"Expected ₹4,200 Cr EV, got {su['enterprise_value']}"

    # PPA Goodwill & DTL check
    ppa = engine.calculate_purchase_price_allocation()
    assert abs(ppa["deferred_tax_liability"] - 151.02) < 0.1, f"Expected DTL ₹151.02 Cr, got {ppa['deferred_tax_liability']}"
    assert abs(ppa["goodwill_created"] - 2351.02) < 0.1, f"Expected Goodwill ₹2,351.02 Cr, got {ppa['goodwill_created']}"

    # Synergies
    assert sum(engine.inputs.synergies_phasing) == 380.0, "Synergies must sum to ₹380 Cr"

    # Year 2 EPS Accretion
    bridge = engine.get_year2_accretion_bridge()
    assert bridge["standalone_eps"] == 45.00, f"Expected Standalone EPS ₹45.00, got {bridge['standalone_eps']}"
    assert abs(bridge["pro_forma_eps"] - 47.88) < 0.05, f"Expected Pro-Forma EPS ₹47.88, got {bridge['pro_forma_eps']}"
    assert abs(bridge["accretion_pct"] - 6.4) < 0.1, f"Expected +6.4% Accretion, got {bridge['accretion_pct']:.2f}%"

    print(f"[TEST PASS] M&A Engine: ₹4,200 Cr EV, ₹380 Cr Synergies, ₹45.00 -> ₹47.88 (+{bridge['accretion_pct']:.1f}% Y2 EPS Accretion).")


def test_debt_structures_and_irr():
    opt = DebtStructuresOptimizer()
    df_structs = opt.evaluate_all_structures()
    assert len(df_structs) == 5, f"Expected 5 structures, got {len(df_structs)}"

    # Structure 4 (Optimal Hybrid)
    struct_4 = df_structs[df_structs["Structure"].str.contains("Optimal Hybrid")].iloc[0]
    assert abs(struct_4["Net Debt / EBITDA"] - 3.08) < 0.02, f"Expected 3.08x Net Debt/EBITDA, got {struct_4['Net Debt / EBITDA']}"
    assert struct_4["Net Debt / EBITDA"] <= 3.20, "Structure 4 must adhere to 3.2x covenant"
    assert struct_4["Covenant Status"] == "PASS (Compliant)"
    assert abs(struct_4["5-Year Project IRR (%)"] - 18.7) < 0.1, f"Expected 18.7% IRR, got {struct_4['5-Year Project IRR (%)']}"

    # Structure 1 (Base Unlevered)
    struct_1 = df_structs[df_structs["Structure"].str.contains("100% Stock")].iloc[0]
    assert abs(struct_1["5-Year Project IRR (%)"] - 14.2) < 0.1, f"Expected 14.2% Base IRR, got {struct_1['5-Year Project IRR (%)']}"

    # FCF Cash sweep
    df_sweep = opt.build_fcf_debt_sweep_schedule()
    assert len(df_sweep) == 5
    assert df_sweep.iloc[-1]["Net Debt / EBITDA"] < 0.5, "Expected deleveraging to <0.5x by Year 5"

    print(f"[TEST PASS] Capital Structures & LBO IRR: Structure #4 Net Debt/EBITDA = {struct_4['Net Debt / EBITDA']}x (<3.2x Covenant), IRR boosted from {struct_1['5-Year Project IRR (%)']}% to {struct_4['5-Year Project IRR (%)']}%.")


def test_excel_workbook():
    if not os.path.exists(EXCEL_OUTPUT_PATH):
        generate_wall_street_excel_model(EXCEL_OUTPUT_PATH)
    assert os.path.exists(EXCEL_OUTPUT_PATH)

    wb = openpyxl.load_workbook(EXCEL_OUTPUT_PATH, data_only=True)
    expected_tabs = [
        "Kaggle_M&A_Target_Screener",
        "Sources_Uses_&_Goodwill_PPA",
        "5Yr_ProForma_IS_&_Synergies",
        "EPS_Accretion_Dilution_Bridge",
        "5_Debt_Structures_&_Cash_Sweep",
        "Sensitivity_Matrices",
    ]
    for tab in expected_tabs:
        assert tab in wb.sheetnames, f"Missing tab: {tab}"
        assert wb[tab].max_row >= 10, f"Tab {tab} has insufficient rows"
    print(f"[TEST PASS] Excel Model: All 6 Wall Street tabs validated in {EXCEL_OUTPUT_PATH}.")


if __name__ == "__main__":
    print("=" * 70)
    print("RUNNING INSTITUTIONAL M&A & LBO VALIDATION TEST SUITE")
    print("=" * 70)
    test_sqlite_database()
    test_sql_screener()
    test_deal_calibration_and_accretion()
    test_debt_structures_and_irr()
    test_excel_workbook()
    print("=" * 70)
    print("ALL TESTS PASSED WITH 100% PRECISION!")
    print("=" * 70)
